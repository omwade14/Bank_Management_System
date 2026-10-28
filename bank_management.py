
import openpyxl
from pathlib import Path

FILE = "bank.xlsx"
HEADERS = ["Name", "Mobile", "Balance"]


# Create Excel file if it does not exist
def initialize_file():
    if not Path(FILE).exists():
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Accounts"
        ws.append(HEADERS)
        wb.save(FILE)
        wb.close()
    else:
        wb = openpyxl.load_workbook(FILE)
        ws = wb.active

        if ws.max_row == 1 and ws["A1"].value is None:
            ws.append(HEADERS)

        wb.save(FILE)
        wb.close()


def get_mobile():
    while True:
        mobile = input("Enter 10-digit Mobile Number: ").strip()

        if mobile.isdigit() and len(mobile) == 10:
            return mobile

        print("Invalid mobile number! Enter 10 digits.")


def get_amount(message, allow_zero=False):
    while True:
        try:
            amount = round(float(input(message)), 2)

            if amount < 0 or (amount == 0 and not allow_zero):
                print("Enter a valid positive amount!")
                continue

            return amount

        except ValueError:
            print("Invalid amount! Please enter a number.")


def find_account(ws, mobile):
    for row in range(2, ws.max_row + 1):
        saved_mobile = ws.cell(row, 2).value

        if saved_mobile is not None:
            if str(saved_mobile).strip() == mobile:
                return row

    return None


# 1. Create New Account
def add_new():
    name = input("Enter Your Name: ").strip()

    if not name:
        print("Name cannot be empty!")
        return

    mobile = get_mobile()
    opening_balance = get_amount(
        "Enter Opening Balance (Rs.): ",
        allow_zero=True
    )

    wb = openpyxl.load_workbook(FILE)
    try:
        ws = wb.active

        if find_account(ws, mobile) is not None:
            print("An account with this mobile number already exists!")
            return

        ws.append([name, mobile, opening_balance])
        wb.save(FILE)

        print("\nAccount Created Successfully!")
        print("Name:", name)
        print("Mobile:", mobile)
        print("Opening Balance: Rs.", f"{opening_balance:.2f}")

    finally:
        wb.close()


# 2. Deposit Money
def deposit_money():
    mobile = get_mobile()
    amount = get_amount("Enter Amount to Deposit (Rs.): ")

    wb = openpyxl.load_workbook(FILE)
    try:
        ws = wb.active
        row = find_account(ws, mobile)

        if row is None:
            print("Account Not Found!")
            return

        balance = float(ws.cell(row, 3).value or 0)
        new_balance = round(balance + amount, 2)

        ws.cell(row, 3).value = new_balance
        wb.save(FILE)

        print("Money Deposited Successfully!")
        print("Current Balance: Rs.", f"{new_balance:.2f}")

    finally:
        wb.close()


# 3. Withdraw Money
def withdrawal_money():
    mobile = get_mobile()
    amount = get_amount("Enter Amount to Withdraw (Rs.): ")

    wb = openpyxl.load_workbook(FILE)
    try:
        ws = wb.active
        row = find_account(ws, mobile)

        if row is None:
            print("Account Not Found!")
            return

        balance = float(ws.cell(row, 3).value or 0)

        if amount > balance:
            print("Insufficient Balance!")
            print("Available Balance: Rs.", f"{balance:.2f}")
            return

        new_balance = round(balance - amount, 2)
        ws.cell(row, 3).value = new_balance
        wb.save(FILE)

        print("Money Withdrawn Successfully!")
        print("Remaining Balance: Rs.", f"{new_balance:.2f}")

    finally:
        wb.close()


# 4. View Balance
def view_balance():
    mobile = get_mobile()

    wb = openpyxl.load_workbook(FILE)
    try:
        ws = wb.active
        row = find_account(ws, mobile)

        if row is None:
            print("Account Not Found!")
            return

        print("\n----- ACCOUNT DETAILS -----")
        print("Name:", ws.cell(row, 1).value)
        print("Mobile:", ws.cell(row, 2).value)
        print("Current Balance: Rs.",
              f"{float(ws.cell(row, 3).value or 0):.2f}")

    finally:
        wb.close()


# 5. Display All Accounts
def display_accounts():
    wb = openpyxl.load_workbook(FILE)
    try:
        ws = wb.active

        if ws.max_row <= 1:
            print("No accounts available!")
            return

        print("\n----- ALL BANK ACCOUNTS -----")
        print(f"{'Name':<20}{'Mobile':<15}{'Balance (Rs.)':>15}")
        print("-" * 50)

        for row in ws.iter_rows(min_row=2, values_only=True):
            name, mobile, balance = row[:3]

            if name is not None:
                print(
                    f"{str(name):<20}"
                    f"{str(mobile):<15}"
                    f"{float(balance or 0):>15.2f}"
                )

    finally:
        wb.close()


# Main Menu
def menu():
    initialize_file()

    while True:
        print("\n========== WELCOME TO THE BANK ==========")
        print("1. Add New Account")
        print("2. Deposit Money")
        print("3. Withdraw Money")
        print("4. View Balance")
        print("5. Display All Accounts")
        print("6. Exit")
        print("=========================================")

        choice = input("Enter Your Choice (1-6): ").strip()

        if choice == "1":
            add_new()
        elif choice == "2":
            deposit_money()
        elif choice == "3":
            withdrawal_money()
        elif choice == "4":
            view_balance()
        elif choice == "5":
            display_accounts()
        elif choice == "6":
            print("Thank You for Using Our Bank System!")
            break
        else:
            print("Invalid Choice! Please select 1 to 6.")


if __name__ == "__main__":
    try:
        menu()
    except (PermissionError, OSError) as error:
        print("File Error:", error)
        print("Close bank.xlsx in Excel and try again.")
